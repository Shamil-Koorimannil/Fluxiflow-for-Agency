import io
import csv
import openpyxl
import datetime

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 5000
MAX_COLS_PER_SHEET = 100
MAX_TOTAL_CELLS = 50000

class SpreadsheetImportError(Exception):
    pass

def col_idx_to_name(idx: int) -> str:
    """Convert 1-based column index to Excel column name (e.g. 1 -> A, 27 -> AA)."""
    name = ''
    while idx > 0:
        idx, remainder = divmod(idx - 1, 26)
        name = chr(65 + remainder) + name
    return name


def import_spreadsheet_file(file_obj, filename: str) -> tuple[dict, list[str]]:
    content = file_obj.read()
    if len(content) > MAX_FILE_SIZE:
        raise SpreadsheetImportError("This spreadsheet is too large for Keep (exceeds maximum file size limit of 10 MB).")

    import_warnings = []
    filename_lower = filename.lower()

    # Magic Bytes & File Format Autodetection
    if content.startswith(b'PK\x03\x04'):
        return parse_xlsx(content, import_warnings)
    elif content.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        return parse_xls(content, import_warnings)
    else:
        if filename_lower.endswith('.xlsx'):
            return parse_xlsx(content, import_warnings)
        elif filename_lower.endswith('.xls'):
            return parse_xls(content, import_warnings)
        else:
            return parse_csv(content, import_warnings)


def parse_xlsx(content: bytes, warnings: list[str]) -> tuple[dict, list[str]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    except Exception:
        raise SpreadsheetImportError("Unable to import this spreadsheet. The file may be damaged or unsupported.")

    if len(wb.sheetnames) > MAX_SHEETS:
        raise SpreadsheetImportError("This spreadsheet is too large for Keep (exceeds maximum sheet limit of 20).")

    sheets_data = []
    total_cells = 0

    for sheet_idx, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        cells_dict = {}
        max_row = min(ws.max_row or 0, MAX_ROWS_PER_SHEET)
        max_col = min(ws.max_column or 0, MAX_COLS_PER_SHEET)

        if ws.max_row and ws.max_row > MAX_ROWS_PER_SHEET:
            warnings.append(f"Sheet '{sheet_name}' rows truncated to {MAX_ROWS_PER_SHEET}.")
        if ws.max_column and ws.max_column > MAX_COLS_PER_SHEET:
            warnings.append(f"Sheet '{sheet_name}' columns truncated to {MAX_COLS_PER_SHEET}.")

        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is None and not cell.comment:
                    continue

                total_cells += 1
                if total_cells > MAX_TOTAL_CELLS:
                    raise SpreadsheetImportError("This spreadsheet is too large for Keep (exceeds maximum cell limit of 50,000).")

                col_name = col_idx_to_name(c)
                cell_key = f"{col_name}{r}"

                val = str(cell.value) if cell.value is not None else ""
                formula_str = ""
                if str(val).startswith('='):
                    formula_str = str(val)
                    val = ""

                fmt = {}
                if cell.font:
                    if cell.font.bold: fmt['bold'] = True
                    if cell.font.italic: fmt['italic'] = True
                    if cell.font.underline: fmt['underline'] = True
                    if cell.font.size: fmt['fontSize'] = int(cell.font.size)
                    if cell.font.color and hasattr(cell.font.color, 'rgb') and cell.font.color.rgb:
                        rgb = str(cell.font.color.rgb)
                        if len(rgb) == 8: rgb = rgb[2:]
                        fmt['color'] = f"#{rgb}"

                if cell.fill and hasattr(cell.fill, 'start_color') and cell.fill.start_color:
                    if hasattr(cell.fill.start_color, 'rgb') and cell.fill.start_color.rgb:
                        rgb = str(cell.fill.start_color.rgb)
                        if len(rgb) == 8: rgb = rgb[2:]
                        if rgb != '00000000':
                            fmt['bgColor'] = f"#{rgb}"

                if cell.alignment and cell.alignment.horizontal:
                    fmt['align'] = cell.alignment.horizontal

                cells_dict[cell_key] = {
                    "value": val,
                    "formula": formula_str,
                    "format": fmt
                }

        sheets_data.append({
            "id": f"sheet_{sheet_idx + 1}",
            "name": sheet_name,
            "frozenRows": 0,
            "frozenCols": 0,
            "cells": cells_dict,
            "rowHeights": {},
            "colWidths": {}
        })

    if getattr(wb, 'vba_archive', None):
        warnings.append("VBA Macros were omitted during import for security.")

    return {"sheets": sheets_data}, warnings


def parse_xls(content: bytes, warnings: list[str]) -> tuple[dict, list[str]]:
    try:
        import xlrd
    except ImportError:
        raise SpreadsheetImportError("XLS support module (xlrd) is missing from python environment.")

    try:
        wb = xlrd.open_workbook(file_contents=content)
    except Exception:
        raise SpreadsheetImportError("Unable to import this legacy spreadsheet (.xls). The file may be damaged or malformed.")

    if wb.nsheets > MAX_SHEETS:
        raise SpreadsheetImportError("This spreadsheet is too large for Keep (exceeds maximum sheet limit of 20).")

    sheets_data = []
    total_cells = 0

    for sheet_idx in range(wb.nsheets):
        ws = wb.sheet_by_index(sheet_idx)
        sheet_name = ws.name
        cells_dict = {}
        max_row = min(ws.nrows, MAX_ROWS_PER_SHEET)
        max_col = min(ws.ncols, MAX_COLS_PER_SHEET)

        if ws.nrows > MAX_ROWS_PER_SHEET:
            warnings.append(f"Sheet '{sheet_name}' rows truncated to {MAX_ROWS_PER_SHEET}.")
        if ws.ncols > MAX_COLS_PER_SHEET:
            warnings.append(f"Sheet '{sheet_name}' columns truncated to {MAX_COLS_PER_SHEET}.")

        for r in range(max_row):
            for c in range(max_col):
                cell_type = ws.cell_type(r, c)
                raw_val = ws.cell_value(r, c)

                if cell_type == xlrd.XL_CELL_EMPTY or raw_val is None or raw_val == "":
                    continue

                total_cells += 1
                if total_cells > MAX_TOTAL_CELLS:
                    raise SpreadsheetImportError("This spreadsheet is too large for Keep (exceeds maximum cell limit of 50,000).")

                col_name = col_idx_to_name(c + 1)
                cell_key = f"{col_name}{r + 1}"

                formatted_val = ""
                if cell_type == xlrd.XL_CELL_DATE:
                    try:
                        dt_tuple = xlrd.xldate_as_tuple(raw_val, wb.datemode)
                        if dt_tuple[0:3] == (0, 0, 0):
                            formatted_val = f"{dt_tuple[3]:02d}:{dt_tuple[4]:02d}:{dt_tuple[5]:02d}"
                        else:
                            formatted_val = f"{dt_tuple[0]:04d}-{dt_tuple[1]:02d}-{dt_tuple[2]:02d}"
                    except Exception:
                        formatted_val = str(raw_val)
                elif cell_type == xlrd.XL_CELL_NUMBER:
                    if isinstance(raw_val, float) and raw_val.is_integer():
                        formatted_val = str(int(raw_val))
                    else:
                        formatted_val = str(raw_val)
                elif cell_type == xlrd.XL_CELL_BOOLEAN:
                    formatted_val = "TRUE" if raw_val else "FALSE"
                elif cell_type == xlrd.XL_CELL_ERROR:
                    warnings.append(f"Sheet '{sheet_name}' cell {cell_key} contained an error.")
                    formatted_val = "#ERROR!"
                else:
                    formatted_val = str(raw_val)

                cells_dict[cell_key] = {
                    "value": formatted_val,
                    "formula": "",
                    "format": {}
                }

        sheets_data.append({
            "id": f"sheet_{sheet_idx + 1}",
            "name": sheet_name,
            "frozenRows": 0,
            "frozenCols": 0,
            "cells": cells_dict,
            "rowHeights": {},
            "colWidths": {}
        })

    return {"sheets": sheets_data}, warnings


def parse_csv(content: bytes, warnings: list[str]) -> tuple[dict, list[str]]:
    try:
        decoded = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            decoded = content.decode('latin-1')
        except Exception:
            raise SpreadsheetImportError("Unable to decode CSV file encoding.")

    reader = csv.reader(io.StringIO(decoded))
    cells_dict = {}
    total_cells = 0

    for r_idx, row in enumerate(reader):
        if r_idx >= MAX_ROWS_PER_SHEET:
            warnings.append(f"CSV rows truncated to {MAX_ROWS_PER_SHEET}.")
            break

        for c_idx, val in enumerate(row):
            if c_idx >= MAX_COLS_PER_SHEET:
                continue
            if not val:
                continue

            total_cells += 1
            if total_cells > MAX_TOTAL_CELLS:
                raise SpreadsheetImportError("This CSV file is too large for Keep (exceeds maximum cell limit of 50,000).")

            col_name = col_idx_to_name(c_idx + 1)
            cell_key = f"{col_name}{r_idx + 1}"

            cells_dict[cell_key] = {
                "value": str(val),
                "formula": "",
                "format": {}
            }

    sheets_data = [{
        "id": "sheet_1",
        "name": "Sheet1",
        "frozenRows": 0,
        "frozenCols": 0,
        "cells": cells_dict,
        "rowHeights": {},
        "colWidths": {}
    }]

    return {"sheets": sheets_data}, warnings
