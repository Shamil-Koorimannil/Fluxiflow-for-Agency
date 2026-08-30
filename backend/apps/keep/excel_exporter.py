import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def parse_cell_key(cell_key: str) -> tuple[int, int]:
    """Convert cell key like 'A1' or 'AB12' into 1-based (row, col) tuple."""
    import re
    match = re.match(r"([A-Za-z]+)(\d+)", cell_key)
    if not match:
        return (1, 1)
    col_str, row_str = match.groups()
    col_str = col_str.upper()
    col_idx = 0
    for char in col_str:
        col_idx = col_idx * 26 + (ord(char) - ord('A') + 1)
    return (int(row_str), col_idx)


def export_to_xlsx(spreadsheet_data: dict) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default active sheet
    wb.remove(wb.active)

    sheets = spreadsheet_data.get('sheets', [])
    if not sheets:
        sheets = [{"id": "sheet_1", "name": "Sheet1", "cells": {}}]

    for sheet_info in sheets:
        ws = wb.create_sheet(title=sheet_info.get('name', 'Sheet'))
        cells = sheet_info.get('cells', {})

        for cell_key, cell_val in cells.items():
            row_idx, col_idx = parse_cell_key(cell_key)
            val = cell_val.get('value', '')
            formula = cell_val.get('formula', '')
            fmt = cell_val.get('format', {})

            cell = ws.cell(row=row_idx, column=col_idx)
            if formula and formula.startswith('='):
                cell.value = formula
            else:
                cell.value = val

            # Apply font styling
            font_kwargs = {}
            if fmt.get('bold'): font_kwargs['bold'] = True
            if fmt.get('italic'): font_kwargs['italic'] = True
            if fmt.get('underline'): font_kwargs['underline'] = 'single'
            if fmt.get('fontSize'): font_kwargs['size'] = fmt.get('fontSize')
            if fmt.get('color'):
                color_hex = fmt['color'].lstrip('#')
                font_kwargs['color'] = f"FF{color_hex}" if len(color_hex) == 6 else color_hex

            if font_kwargs:
                cell.font = Font(**font_kwargs)

            # Apply background color
            if fmt.get('bgColor'):
                bg_hex = fmt['bgColor'].lstrip('#')
                fill_color = f"FF{bg_hex}" if len(bg_hex) == 6 else bg_hex
                cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')

            # Apply alignment
            if fmt.get('align'):
                cell.alignment = Alignment(horizontal=fmt['align'])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_to_csv(spreadsheet_data: dict, sheet_id: str | None = None) -> bytes:
    sheets = spreadsheet_data.get('sheets', [])
    target_sheet = None

    if sheet_id:
        for s in sheets:
            if s.get('id') == sheet_id:
                target_sheet = s
                break

    if not target_sheet and sheets:
        target_sheet = sheets[0]

    if not target_sheet:
        return b""

    cells = target_sheet.get('cells', {})
    if not cells:
        return b""

    # Find max row and column
    max_row = 1
    max_col = 1
    for cell_key in cells.keys():
        r, c = parse_cell_key(cell_key)
        if r > max_row: max_row = r
        if c > max_col: max_col = c

    # Build grid matrix
    grid = [["" for _ in range(max_col)] for _ in range(max_row)]
    for cell_key, cell_val in cells.items():
        r, c = parse_cell_key(cell_key)
        val = cell_val.get('value', '')
        formula = cell_val.get('formula', '')
        grid[r - 1][c - 1] = formula if formula else val

    out = io.StringIO()
    writer = csv.writer(out)
    for row in grid:
        writer.writerow(row)

    return out.getvalue().encode('utf-8-sig')
